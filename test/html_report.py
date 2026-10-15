"""Offline HTML reports; also convert an existing DocPanel Excel report."""
# Allow direct execution from test/ while keeping configuration at project root.
import sys as _sys
from pathlib import Path as _Path
for _directory in (_Path(__file__).resolve().parent, _Path(__file__).resolve().parent.parent):
    if str(_directory) not in _sys.path:
        _sys.path.insert(0, str(_directory))

import argparse
import base64
from collections import Counter
from datetime import datetime
from html import escape
import os
from pathlib import Path
from urllib.parse import quote

HEADERS = ["Run time", "Section", "Check", "Count", "Result", "Date range",
           "Details", "API Count", "Unit"]
JOB_HEADERS = ["Run time", "Job type", "Status", "Count", "Sample result",
               "Job ID", "Date range", "Job JSON", "Details"]


def text(value):
    return escape("" if value is None else str(value), quote=True)


def embedded_screenshot(path):
    path = Path(path)
    if not path.is_file():
        return "<p>Screenshot unavailable: " + text(path.name) + "</p>"
    mime = "image/png" if path.suffix.lower() == ".png" else "image/jpeg"
    data = "data:" + mime + ";base64," + base64.b64encode(path.read_bytes()).decode("ascii")
    return ('<figure><figcaption>' + text(path.stem.replace("_", " "))
            + '</figcaption><img class="screenshot" alt="Captured application screenshot" src="'
            + data + '"></figure>')


LOAD_SECTIONS = {"Case Details", "Unity Viewer", "Eunity Viewer", "PowerScribe"}


def load_label(result):
    return {"PASS": "Loaded", "FAIL": "Not loaded", "ERROR": "Not loaded"}.get(result, "Not checked")


def cell(key, value, output, section=None):
    if key == "Details" and isinstance(value, str) and "Screenshot: " in value:
        prefix, candidate = value.rsplit("Screenshot: ", 1)
        screenshot = Path(candidate.strip())
        if screenshot.is_file() and screenshot.suffix.lower() in (".png", ".jpg", ".jpeg"):
            return text(prefix) + embedded_screenshot(screenshot)
        return text(prefix) + '<p>Screenshot unavailable: ' + text(candidate) + '</p>'
    if key == "Status" and section == "Zendesk Unassigned Tickets":
        status = str(value or "").strip().lower()
        colors = {"open":("#d93042", "#ffffff"), "new":("#ffa340", "#38220d")}
        if status in colors:
            background, foreground = colors[status]
            return ('<span class="ticket-status" style="display:inline-block;min-width:48px;text-align:center;'
                    'padding:2px 9px;border-radius:4px;font-size:12px;font-weight:600;background:'
                    + background + ';color:' + foreground + '">' + text(status.title()) + '</span>')
    if key == "Ticket URL" and value:
        from urllib.parse import urlsplit
        parsed = urlsplit(str(value))
        if parsed.scheme == "https" and parsed.hostname and parsed.hostname.endswith(".zendesk.com") and parsed.path.startswith("/agent/tickets/"):
            return '<a href="' + text(value) + '" target="_blank" rel="noopener noreferrer">Open ticket</a>'
    if key == "Job JSON" and value:
        return "<details><summary>View JSON</summary><pre>" + text(value) + "</pre></details>"
    if key in ("Result", "Sample result", "Job sample"):
        style = {"PASS": "pass", "CAPTURED": "pass", "FAIL": "fail",
                 "ERROR": "fail", "SKIP": "pending", "UNVERIFIED": "pending"}.get(str(value), "neutral")
        label = load_label(value) if key == "Result" and section in LOAD_SECTIONS else value
        return '<span class="badge ' + style + '">' + text(label) + "</span>"
    return text(value)


def table(rows, headers, output, section=None):
    if not rows:
        return '<p class="empty">No results recorded.</p>'
    # Zero is meaningful; only None and whitespace-only values are empty.
    headers = [h for h in headers if any(
        row.get(h) is not None and str(row.get(h)).strip() != ""
        for row in rows
    )]
    heading = "".join("<th scope=col>" + text(h) + "</th>" for h in headers)
    body = "".join("<tr>" + "".join("<td>" + cell(h, row.get(h), output, row.get("Section", section)) + "</td>"
                                   for h in headers) + "</tr>" for row in rows)
    return '<div class="table-wrap"><table><thead><tr>' + heading + "</tr></thead><tbody>" + body + "</tbody></table></div>"


def merge_job_samples(checks, jobs):
    """Attach each captured sample to its count row without changing the count."""
    merged = [dict(row) for row in checks]
    for sample in jobs:
        section = str(sample.get("Job type", "")) + " jobs"
        row = next((row for row in merged
                    if row.get("Section") == section
                    and row.get("Check") == sample.get("Status")
                    and row.get("Date range", "") == sample.get("Date range", "")
                    and "Job sample" not in row), None)
        if row is None:
            row = {"Run time":sample.get("Run time"), "Section":section,
                   "Check":sample.get("Status"), "Count":sample.get("Count"),
                   "Result":"COUNT", "Date range":sample.get("Date range")}
            merged.append(row)
        row["Job sample"] = sample.get("Sample result")
        row["Job ID"] = sample.get("Job ID")
        row["Job JSON"] = sample.get("Job JSON")
        row["Sample details"] = sample.get("Details")
    return merged


def jobs_indication(rows):
    expected = {(job, status) for job in ("SR Mach7", "Study Routing Mach7", "Case Dicom CMove")
                for status in ("NEW", "INPROGRESS")}
    selected = [row for row in rows if (row.get("Job type"), row.get("Status")) in expected]
    keys = {(row.get("Job type"), row.get("Status")) for row in selected}
    if (keys != expected or len(selected) != len(expected)
            or any(row.get("Result") != "COUNT" or type(row.get("Count")) not in (int, float)
                   or row["Count"] < 0 for row in selected)):
        return "NEW / INPROGRESS counts are incomplete; see the check results."
    periods = {row.get("Date range", "") for row in selected}
    if len(periods) != 1 or not next(iter(periods)):
        return "NEW / INPROGRESS counts have different or missing date ranges; see details."
    period = next(iter(periods))
    wording = "the reported date range"
    try:
        start, end = period.split(" to ")
        if (datetime.fromisoformat(end) - datetime.fromisoformat(start)).days == 2:
            wording = "the last three days"
    except ValueError:
        pass
    new = sum(row["Count"] for row in selected if row["Status"] == "NEW")
    progress = sum(row["Count"] for row in selected if row["Status"] == "INPROGRESS")
    if new == 0 and progress == 0:
        return f"No jobs are in NEW or INPROGRESS across the three monitored job types for {wording} ({period})."
    return (f"Jobs in NEW: {new}; INPROGRESS: {progress} across the three monitored job types "
            f"for {wording} ({period}).")


def jobs_status(rows):
    monitored = [row for row in rows if row.get("Status") in ("NEW", "INPROGRESS")]
    if any(type(row.get("Count")) in (int, float) and row["Count"] > 0 for row in monitored):
        return "FAIL"
    return "PASS" if jobs_indication(rows).startswith("No jobs are") else ""


def save_html_report(results, output_path=None, excel_path=None, screenshot_paths=()):
    rows = list(results)
    output = Path(output_path) if output_path else Path(__file__).resolve().parent.parent / "reports" / (
        f"docpanel_report_{datetime.now():%Y%m%d_%H%M%S_%f}.html")
    output.parent.mkdir(parents=True, exist_ok=True)
    checks = [r for r in rows if r.get("_kind") not in ("job_sample", "db_table")]
    datasets = [r for r in rows if r.get("_kind") == "db_table"]
    jobs = [r for r in rows if r.get("_kind") == "job_sample"]
    checks = merge_job_samples(checks, jobs)
    for row in checks:
        if row.get("Section") == "Unity Viewer":
            row["Section"] = "Eunity Viewer"
        section = row.get("Section", "")
        if (section in ("Report Pending", "Report Failures") and row.get("Check") == "Query result rows"
                and row.get("Result") in ("COUNT", "PASS", "FAIL")
                and type(row.get("Count")) in (int, float) and row["Count"] >= 0):
            row["Result"] = "FAIL" if row["Count"] > 50 else "PASS"
        if section in ("SR Mach7 jobs", "Study Routing Mach7 jobs", "Case Dicom CMove jobs"):
            row["Job type"] = section.removesuffix(" jobs")
            row["Status"] = row.pop("Check", "")
            row["Section"] = "Jobs Dashboard"
    counts = Counter(str(r.get("Result", "")) for r in checks)
    cards = "".join('<div class="card"><strong>' + str(number) + '</strong><span>' + label + "</span></div>"
                    for label, number in [("Recorded checks", len(checks)),
                                          ("PASS", counts["PASS"]), ("FAIL", counts["FAIL"]),
                                          ("COUNT", counts["COUNT"]),
                                          ("SKIP / UNVERIFIED", counts["SKIP"] + counts["UNVERIFIED"])])
    excel_link = ""
    if excel_path:
        href = quote(Path(os.path.relpath(Path(excel_path).resolve(), output.parent.resolve())).as_posix(), safe="/")
        excel_link = '<a class="download" href="' + text(href) + '">Open Excel report</a>'
    sections = ""
    attached_datasets = set()
    dialogs = ""
    extra_images = {}
    for path in screenshot_paths:
        name = Path(path).name
        section_name = ("Case Assignment" if name.startswith("case_assignment_")
                        else "Legacy Operations" if name.startswith("legacy_operations_")
                        else "Screenshots")
        extra_images.setdefault(section_name, []).append(Path(path))
    for index, section in enumerate(dict.fromkeys(r.get("Section", "") for r in checks)):
        section_rows = [dict(r) for r in checks if r.get("Section", "") == section]
        images = list(extra_images.pop(section, []))
        for row in section_rows:
            details = row.get("Details")
            if isinstance(details, str) and "Screenshot: " in details:
                prefix, candidate = details.rsplit("Screenshot: ", 1)
                images.append(Path(candidate.strip()))
                row["Details"] = prefix.strip()
        images = list(dict.fromkeys(path.resolve() for path in images))
        link = ""
        if images:
            dialog_id = f"screenshots-{index}"
            link = (' <a class="screenshot-link" href="#' + dialog_id
                    + '" data-dialog="' + dialog_id + '">View screenshots (' + str(len(images)) + ')</a>')
            dialogs += ('<dialog id="' + dialog_id + '" aria-labelledby="' + dialog_id + '-title">'
                        + '<form method="dialog" class="dialog-close-bar"><button class="close" type="submit" aria-label="Close screenshots" title="Close screenshots">&#215;</button></form>'
                        + '<h2 id="' + dialog_id + '-title">' + text(section) + ' screenshots</h2>'
                        + "".join(embedded_screenshot(path) for path in images) + '</dialog>')
        verdicts = [row.get("Result") for row in section_rows if row.get("Result") != "COUNT"]
        status = ("FAIL" if any(value in ("FAIL", "ERROR") for value in verdicts)
                  else "PASS" if verdicts and all(value == "PASS" for value in verdicts) else "")
        if section == "Jobs Dashboard":
            status = "FAIL" if status == "FAIL" else jobs_status(section_rows)
        is_load_check = section in LOAD_SECTIONS
        is_tickets = section == "Zendesk Unassigned Tickets"
        collapse = bool(status) or is_load_check or is_tickets
        toggle_label = load_label(status) if is_load_check else status
        toggle_style = status.lower() if status else "pending"
        if status == "FAIL" and not is_load_check:
            toggle_label = "View report"
        if is_tickets and not status:
            toggle_label, toggle_style = "View tickets", "neutral"
        if section == "Jobs Dashboard":
            collapse = True
            toggle_label, toggle_style = "View job details", "neutral"
        panel_id = f"section-details-{index}"
        toggle = (' <button type="button" class="badge ' + toggle_style + ' pass-toggle" data-pass-toggle '
                  + 'aria-expanded="false" aria-controls="' + panel_id + '">' + toggle_label + ' &#9662;</button>') if collapse else ""
        loading_counts = [row["Count"] for row in section_rows
                          if section == "Admin Dashboard" and row.get("Check") == "Loading cases"
                          and isinstance(row.get("Count"), (int, float)) and row["Count"] > 5]
        indication = (' <span class="loading-indication">' + text(loading_counts[0])
                      + ' cases are in Loading status.</span>' if loading_counts else "")
        if section == "Jobs Dashboard":
            indication = ' <span class="jobs-indication">' + text(jobs_indication(section_rows)) + '</span>'
        if section in ("Report Pending", "Report Failures"):
            pending_counts = [row["Count"] for row in section_rows
                              if row.get("Check") == "Query result rows"
                              and type(row.get("Count")) in (int, float) and row["Count"] >= 0]
            if len(pending_counts) == 1:
                indication = (' <span class="jobs-indication ' + status.lower() + '-text">'
                              + text(pending_counts[0]) + (' report(s) pending.' if section == "Report Pending" else ' report failure row(s).') + '</span>')
        if is_tickets:
            count = next((row.get("Count") for row in section_rows if row.get("Count") is not None), None)
            if count is not None:
                indication = ' <span class="jobs-indication">' + text(count) + ' unassigned-view tickets.</span>'
        heading = {"Case Details": "Can you open up cases in DocPanel? - Case Details",
                   "Eunity Viewer": "Images loading - Eunity",
                   "PowerScribe": "Does PowerScribe load?"}.get(section, section or "Results")
        sections += "<section><h2>" + text(heading) + link + toggle + indication + "</h2>"
        headers = (["Run time", "Job type", "Status", "Count", "Result", "Date range",
                    "Job sample", "Job ID", "Job JSON", "Details", "Sample details"]
                   if section == "Jobs Dashboard" else HEADERS)
        content = table(section_rows, headers, output)
        if section in ("Report Pending", "Report Failures", "Zendesk Unassigned Tickets"):
            for dataset_index, dataset in enumerate(datasets):
                if dataset["Section"] == section:
                    content += '<h3>' + ('Pending report details' if section == "Report Pending" else 'Ticket details' if is_tickets else 'Report failure details') + '</h3>'
                    content += (table(dataset["rows"], dataset["columns"], output, section=dataset["Section"])
                                if dataset["rows"] else '<p>No matching records (0 rows).</p>')
                    attached_datasets.add(dataset_index)
        if collapse:
            content = '<div class="section-details" id="' + panel_id + '" hidden>' + content + '</div>'
        sections += content + "</section>"
    if not checks:
        sections = "<p>No checks recorded.</p>"
    for dataset_index, dataset in enumerate(datasets):
        if dataset_index in attached_datasets:
            continue
        sections += "<section><h2>" + text(dataset["Section"]) + " - query results</h2>"
        if dataset["rows"]:
            sections += table(dataset["rows"], dataset["columns"], output, section=dataset["Section"])
        else:
            sections += "<p>No matching records (0 rows).</p>"
        sections += "</section>"
    for name, images in extra_images.items():
        sections += "<section><h2>" + text(name) + "</h2>" + "".join(embedded_screenshot(p) for p in images) + "</section>"
    sections += dialogs
    document = """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>DocPanel automation report</title>
<style>
.jobs-indication.fail-text{color:#b42318}.jobs-indication.pass-text{color:#155d37}
.jobs-indication{display:inline-block;margin-left:12px;font-size:14px;font-weight:500;color:#304b69}
.loading-indication{display:inline-block;margin-left:12px;font-size:14px;font-weight:500;color:#94261e}
.pass-toggle{border:0;cursor:pointer;font-size:13px;margin-left:14px}.pass-toggle:focus-visible{outline:3px solid #126788;outline-offset:2px}.section-details{margin-top:12px}@media print{.section-details[hidden]{display:block}}
dialog{width:92vw;max-width:1600px;max-height:90vh;overflow:auto;border:0;border-radius:12px;padding:24px}dialog::backdrop{background:#0009}.dialog-close-bar{position:sticky;top:0;z-index:10;display:flex;justify-content:flex-end;margin:0;background:#fff}.close{display:grid;place-items:center;width:44px;height:44px;padding:0;border:1px solid #dce4ea;border-radius:50%;background:#fff;color:#182b3b;font-size:30px;line-height:1;cursor:pointer}.close:hover{background:#edf2f6}.close:focus-visible{outline:3px solid #126788;outline-offset:2px}.screenshot-link{font-size:14px;font-weight:normal;margin-left:14px}figure{margin:14px 0}figcaption{font-size:12px;color:#526779;margin-bottom:8px}.screenshot{display:block;max-width:100%;height:auto;border:1px solid #dce4ea}*{box-sizing:border-box}body{margin:0;background:#f3f6fa;color:#182b3b;font:14px/1.5 system-ui,Arial,sans-serif}
header{background:#164b5e;color:white;padding:30px 32px}h1{margin:0 0 8px;font-size:26px}header p{margin:0;color:#d4e8ec}
main{padding:24px 32px}.cards{display:flex;flex-wrap:wrap;gap:14px;margin-bottom:26px}
.card{background:white;border:1px solid #dce4ea;border-radius:10px;padding:16px 24px;min-width:150px}
.card strong{display:block;font-size:27px}.card span{color:#526779}section{margin:28px 0}h2{font-size:19px}
.table-wrap{overflow:auto;background:white;border:1px solid #dce4ea;border-radius:8px}
table{border-collapse:collapse;width:100%}th,td{text-align:left;padding:12px;border-bottom:1px solid #e3eaf0;vertical-align:top}
th{background:#eaf1f5;white-space:nowrap}td{min-width:90px;max-width:550px;overflow-wrap:anywhere}
tr:nth-child(even){background:#fafcfd}.badge{display:inline-block;padding:3px 9px;border-radius:20px;font-weight:600;white-space:nowrap}
.pass{background:#dff3e7;color:#155d37}.fail{background:#ffe2e0;color:#94261e}.pending{background:#fff0c9;color:#735108}
.neutral{background:#e7eef6;color:#304b69}a{color:#126788}.download{display:inline-block;margin-top:15px;color:white}
pre{white-space:pre-wrap;overflow-wrap:anywhere;max-height:400px;overflow:auto;min-width:300px;font:12px/1.5 Consolas,monospace}
summary{cursor:pointer;color:#126788}footer{color:#526779;margin-top:24px}
@media print{header{background:white;color:black}header p{color:black}.table-wrap{overflow:visible}main{padding:0}th,td{padding:5px;font-size:10px}}
</style></head><body>"""
    document += "<header><h1>DocPanel automation report</h1><p>Generated " + text(datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    document += " ? Counts and check outcomes are reported separately.</p>" + excel_link + "</header>"
    document += '<main><div class="cards">' + cards + "</div>" + sections
    document += "<footer>Screenshots are embedded in this HTML file. Job JSON is expandable.</footer></main><script>\ndocument.querySelectorAll('[data-dialog]').forEach(link => {\n    link.addEventListener('click', event => {\n        event.preventDefault();\n        document.getElementById(link.dataset.dialog).showModal();\n    });\n});\n\ndocument.querySelectorAll('[data-pass-toggle]').forEach(button => {\n    button.addEventListener('click', () => {\n        const panel = document.getElementById(button.getAttribute('aria-controls'));\n        const expanded = button.getAttribute('aria-expanded') !== 'true';\n        panel.hidden = !expanded;\n        button.setAttribute('aria-expanded', String(expanded));\n    });\n});\n</script></body></html>"
    from report_theme import apply_theme
    document = apply_theme(document)
    output.write_text(document, encoding="utf-8")
    print(f"HTML report saved: {output.resolve()}", flush=True)
    return output


def convert_excel(excel_path, screenshot_paths=()):
    from openpyxl import load_workbook
    path = Path(excel_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, ())
            if sheet.title in ("Report Pending", "Report Failures", "Zendesk Unassigned Tickets"):
                records.append({"_kind":"db_table", "Section":sheet.title, "columns":list(headers),
                                "rows":[dict(zip(headers, values)) for values in iterator
                                        if any(value is not None for value in values)]})
                continue
            for values in iterator:
                row = dict(zip(headers, values))
                if not any(value is not None for value in values):
                    continue
                if sheet.title == "Job Details":
                    row["_kind"] = "job_sample"
                    # Reassemble chunks created by the Excel export.
                    if str(row.get("JSON part", "1/1")).startswith("1/"):
                        records.append(row)
                    elif records and records[-1].get("_kind") == "job_sample":
                        records[-1]["Job JSON"] = (records[-1].get("Job JSON") or "") + (row.get("Job JSON") or "")
                    else:
                        raise ValueError("Job JSON continuation has no first part.")
                else:
                    records.append(row)
    finally:
        workbook.close()
    return save_html_report(records, path.with_suffix(".html"), path, screenshot_paths)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("excel_file", help="Existing DocPanel .xlsx report")
    convert_excel(parser.parse_args().excel_file)

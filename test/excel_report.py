# Allow direct execution from test/ while keeping configuration at project root.
import sys as _sys
from pathlib import Path as _Path
for _directory in (_Path(__file__).resolve().parent, _Path(__file__).resolve().parent.parent):
    if str(_directory) not in _sys.path:
        _sys.path.insert(0, str(_directory))

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


def record_result(results, section, check, count=None, result="COUNT", date_range="", details="", api_count=None, unit=""):
    if results is not None:
        results.append({
            "Run time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Section": section,
            "Check": check,
            "Count": count,
            "Result": result,
            "Date range": date_range,
            "Details": details,
            "API Count": api_count,
            "Unit": unit,
        })



def record_job_sample(results, job_type, status, count, date_range, result,
                      job_id="", job_json="", details=""):
    if results is not None:
        results.append({
            "_kind": "job_sample",
            "Run time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "Job type": job_type, "Status": status, "Count": count,
            "Sample result": result, "Job ID": job_id, "Date range": date_range,
            "Job JSON": job_json, "Details": details,
        })


def add_job_details_sheet(workbook, samples):
    sheet = workbook.create_sheet("Job Details")
    headers = ["Run time", "Job type", "Status", "Count", "Sample result",
               "Job ID", "Date range", "JSON part", "Job JSON", "Details"]
    sheet.append(headers)
    for sample in samples:
        # Excel cells hold at most 32,767 characters. ASCII JSON chunks avoid
        # truncation and UTF-16 character limits. Join parts in order if split.
        payload = sample["Job JSON"]
        chunks = [payload[i:i + 30000] for i in range(0, len(payload), 30000)] or [""]
        for index, chunk in enumerate(chunks, start=1):
            row = dict(sample, **{"JSON part": f"{index}/{len(chunks)}", "Job JSON": chunk})
            sheet.append([row.get(key) for key in headers])
            sheet.row_dimensions[sheet.max_row].height = 90 if payload else 30
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164B5E")
    for column, width in zip("ABCDEFGHIJ", (22, 26, 20, 12, 22, 40, 28, 12, 100, 65)):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = "#,##0"
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(cell.value, str):
                cell.data_type = "s"
        row[8].font = Font(name="Consolas", size=10)
        if row[4].value in ("CAPTURED", "ERROR"):
            row[4].fill = PatternFill("solid", fgColor="D9EAD3" if row[4].value == "CAPTURED" else "F4CCCC")
    sheet.freeze_panes = "A2"
    table = Table(displayName="JobDetails", ref=f"A1:J{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)



def style_ticket_statuses(sheet):
    headers = [cell.value for cell in sheet[1]]
    if "Status" not in headers:
        return
    index = headers.index("Status")
    colors = {"open":("D93042", "FFFFFF"), "new":("FFA340", "38220D")}
    for row in sheet.iter_rows(min_row=2):
        cell = row[index]
        status = str(cell.value or "").strip().lower()
        if status in colors:
            background, foreground = colors[status]
            cell.fill = PatternFill("solid", fgColor=background)
            cell.font = Font(color=foreground, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def add_database_sheet(workbook, dataset):
    from openpyxl.utils import get_column_letter
    sheet = workbook.create_sheet(dataset["Section"])
    columns = dataset["columns"]
    sheet.append(columns)
    for item in dataset["rows"]:
        sheet.append([item.get(column) for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164B5E")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index in range(1, len(columns) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if dataset["Section"] == "Zendesk Unassigned Tickets":
        style_ticket_statuses(sheet)


def save_excel_report(results, output_dir=None):
    all_results = list(results)
    datasets = [row for row in all_results if row.get("_kind") == "db_table"]
    results = [row for row in all_results if row.get("_kind") != "db_table"]
    samples = [row for row in results if row.get("_kind") == "job_sample"]
    results = [row for row in results if row.get("_kind") != "job_sample"]
    output_dir = Path(output_dir) if output_dir else Path(__file__).resolve().parent.parent / "reports"
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / f"docpanel_report_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DocPanel Results"
    headers = ["Run time", "Section", "Check", "Count", "Result", "Date range", "Details", "API Count", "Unit"]
    sheet.append(headers)
    for row in results:
        sheet.append([row.get(key) for key in headers])
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164B5E")
    for col, width in zip("ABCDEFGHI", (22, 25, 44, 14, 14, 28, 65, 14, 16)):
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows(min_row=2):
        row[3].number_format = "#,##0.##"
        row[7].number_format = "#,##0.##"
        row[6].alignment = Alignment(wrap_text=True, vertical="top")
        status = row[4]
        if status.value in ("PASS", "FAIL"):
            status.fill = PatternFill("solid", fgColor="D9EAD3" if status.value == "PASS" else "F4CCCC")
        # Treat any diagnostic text as literal text, never as an Excel formula.
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.data_type = "s"
    if results:
        table = Table(displayName="DocPanelResults", ref=f"A1:I{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    if samples:
        add_job_details_sheet(workbook, samples)
    for dataset in datasets:
        add_database_sheet(workbook, dataset)
    workbook.save(path)
    print(f"Excel report saved: {path.resolve()}", flush=True)
    from html_report import save_html_report
    save_html_report(all_results, path.with_suffix(".html"), path)
    return path
